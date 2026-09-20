from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT

import pandas as pd


ROOT = Path(__file__).resolve().parents[1]

FIG_DIR = ROOT / "reports" / "figures"

REPORT_FILE = (
    ROOT /
    "reports" /
    "credit_scoring_report.docx"
)

MODEL_RESULTS = (
    ROOT /
    "reports" /
    "model_comparison.csv"
)


# ============================================================
# DOCUMENT
# ============================================================

doc = Document()

section = doc.sections[0]

section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)


# ============================================================
# TITLE
# ============================================================

title = doc.add_paragraph()

title.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = title.add_run(
    "CREDIT SCORING MODEL\n"
)

run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(
    16,
    42,
    67
)

subtitle = title.add_run(
    "End-to-End Machine Learning Project"
)

subtitle.font.size = Pt(15)


doc.add_paragraph(
    "Predicting creditworthiness using applicant financial "
    "and credit-history characteristics."
)


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================

doc.add_heading(
    "1. Executive Summary",
    level=1
)

doc.add_paragraph(
    "This project develops an end-to-end machine-learning "
    "credit scoring system using the UCI Statlog German "
    "Credit dataset. The objective is to classify applicants "
    "as good or bad credit risks while emphasizing model "
    "interpretability, cost-sensitive evaluation and "
    "professional visualization."
)


# ============================================================
# DATASET
# ============================================================

doc.add_heading(
    "2. Dataset",
    level=1
)

doc.add_paragraph(
    "The dataset contains 1,000 applicants and 20 predictor "
    "variables. Variables describe checking accounts, credit "
    "history, loan duration, credit amount, savings, employment, "
    "installment burden, property, housing, existing credits "
    "and other applicant characteristics."
)

doc.add_paragraph(
    "Important limitation: the dataset does not contain a "
    "literal annual-income field. Therefore, income is not "
    "invented. Financial capacity is represented using the "
    "available credit amount, installment rate, checking/savings "
    "status and related variables."
)


# ============================================================
# MODEL RESULTS
# ============================================================

doc.add_heading(
    "3. Model Performance",
    level=1
)

if MODEL_RESULTS.exists():

    results = pd.read_csv(
        MODEL_RESULTS
    )

    table = doc.add_table(
        rows=1,
        cols=len(results.columns)
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    for i, column in enumerate(
        results.columns
    ):
        table.rows[0].cells[i].text = column

    for _, row in results.iterrows():

        cells = table.add_row().cells

        for i, value in enumerate(
            row
        ):

            if isinstance(
                value,
                float
            ):

                cells[i].text = (
                    f"{value:.4f}"
                )

            else:

                cells[i].text = str(
                    value
                )


# ============================================================
# VISUALIZATIONS
# ============================================================

doc.add_heading(
    "4. Exploratory Analysis",
    level=1
)

figures = [
    "01_target_distribution.png",
    "02_credit_amount_risk.png",
    "03_age_distribution.png",
    "04_loan_duration.png",
    "06_credit_history_risk.png",
    "07_correlation_matrix.png",
    "09_confusion_matrix.png",
    "10_roc_curve.png",
    "11_precision_recall_curve.png",
    "12_calibration_curve.png",
    "13_shap_feature_importance.png"
]


for figure in figures:

    path = FIG_DIR / figure

    if path.exists():

        doc.add_picture(
            str(path),
            width=Inches(6.3)
        )

        paragraph = doc.paragraphs[-1]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )


# ============================================================
# CONCLUSION
# ============================================================

doc.add_heading(
    "5. Conclusion",
    level=1
)

doc.add_paragraph(
    "The resulting system provides a complete credit-risk "
    "workflow covering data acquisition, validation, "
    "exploratory analysis, feature engineering, model "
    "development, evaluation, threshold optimization, "
    "explainability and applicant-level prediction."
)

doc.add_paragraph(
    "For production deployment, additional validation on "
    "recent institution-specific data, monitoring for "
    "population drift, probability calibration, fairness "
    "assessment, governance review and regulatory compliance "
    "would be required."
)


# ============================================================
# SAVE
# ============================================================

REPORT_FILE.parent.mkdir(
    parents=True,
    exist_ok=True
)

doc.save(
    REPORT_FILE
)

print(
    f"Report created: {REPORT_FILE}"
)
